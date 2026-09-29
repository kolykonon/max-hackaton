"""Заявление работодателю на доп. день отдыха после донации (ст. 186 ТК РФ).

ФИО берётся из личных данных. Реквизиты работодателя не храним: приходят
с фронта при каждой генерации, пустые поля остаются линиями для заполнения от руки.
"""

import datetime as dt
import io
from dataclasses import dataclass
from enum import StrEnum
from pathlib import Path

from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from fpdf import FPDF

from app.models import Donation, PersonalData, User

FONTS_DIR = Path(__file__).resolve().parent.parent / "assets" / "fonts"
BLANK = "_" * 28
BLANK_DATE = "«___» ____________ 20___ г."

MONTHS_GEN = [
    "января",
    "февраля",
    "марта",
    "апреля",
    "мая",
    "июня",
    "июля",
    "августа",
    "сентября",
    "октября",
    "ноября",
    "декабря",
]


class DocFormat(StrEnum):
    PDF = "pdf"
    DOCX = "docx"


MIME = {
    DocFormat.PDF: "application/pdf",
    DocFormat.DOCX: "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


@dataclass(frozen=True)
class EmployerFields:
    employer_name: str | None = None
    head_position: str | None = None
    head_name: str | None = None
    employee_position: str | None = None
    rest_date: dt.date | None = None
    attach_to_vacation: bool = False


@dataclass(frozen=True)
class LeaveApplication:
    header: list[str]
    title: str
    body: list[str]
    footer_date: str
    footer_sign: str


def ru_date(d: dt.date) -> str:
    return f"«{d.day:02d}» {MONTHS_GEN[d.month - 1]} {d.year} г."


def _full_name(user: User, pd: PersonalData | None) -> str:
    if pd and pd.last_name and pd.first_name:
        return " ".join(p for p in (pd.last_name, pd.first_name, pd.middle_name) if p)
    return " ".join(p for p in (user.last_name, user.first_name) if p)


def _short_name(user: User, pd: PersonalData | None) -> str:
    """Иванов И. И."""
    last = (pd.last_name if pd else None) or user.last_name or ""
    first = (pd.first_name if pd else None) or user.first_name or ""
    middle = pd.middle_name if pd else None
    initials = "".join(f"{p[0]}. " for p in (first, middle) if p).strip()
    return f"{last} {initials}".strip()


def build_application(
    user: User, pd: PersonalData | None, donation: Donation, fields: EmployerFields
) -> LeaveApplication:
    header = [
        fields.head_position or BLANK,
        fields.employer_name or BLANK,
        fields.head_name or BLANK,
        "",
        f"от работника: {_full_name(user, pd) or BLANK}",
        f"должность: {fields.employee_position or '_' * 17}",
    ]
    donated = ru_date(donation.donated_on)
    if fields.attach_to_vacation:
        request = (
            "В соответствии с частью 4 статьи 186 Трудового кодекса Российской "
            "Федерации прошу присоединить дополнительный день отдыха, "
            f"предоставляемый в связи со сдачей крови и её компонентов {donated}, "
            "к ежегодному оплачиваемому отпуску."
        )
    else:
        rest = ru_date(fields.rest_date) if fields.rest_date else BLANK_DATE
        request = (
            "В соответствии с частью 4 статьи 186 Трудового кодекса Российской "
            f"Федерации прошу предоставить мне дополнительный день отдыха {rest} "
            f"в связи со сдачей крови и её компонентов {donated} "
            "с сохранением среднего заработка."
        )
    body = [
        request,
        f"Приложение: справка о донации (форма № 402/у) от {donated}",
    ]
    return LeaveApplication(
        header=header,
        title="ЗАЯВЛЕНИЕ",
        body=body,
        footer_date=BLANK_DATE,
        footer_sign=f"____________ / {_short_name(user, pd)}",
    )


def render_pdf(app: LeaveApplication) -> bytes:
    pdf = FPDF(format="A4")
    pdf.set_margins(left=25, top=20, right=15)
    pdf.add_font("Serif", "", str(FONTS_DIR / "LiberationSerif-Regular.ttf"))
    pdf.add_font("Serif", "B", str(FONTS_DIR / "LiberationSerif-Bold.ttf"))
    pdf.add_page()
    pdf.set_font("Serif", size=14)

    right_x, right_w = 110, pdf.w - 110 - pdf.r_margin
    for line in app.header:
        pdf.set_x(right_x)
        pdf.multi_cell(right_w, 7, line, new_x="LMARGIN", new_y="NEXT")

    pdf.ln(18)
    pdf.set_font("Serif", "B", 14)
    pdf.cell(0, 8, app.title, align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(8)
    pdf.set_font("Serif", size=14)
    for paragraph in app.body:
        pdf.multi_cell(
            0, 7.5, "        " + paragraph, align="J", new_x="LMARGIN", new_y="NEXT"
        )
        pdf.ln(4)

    pdf.ln(16)
    y = pdf.get_y()
    pdf.cell(80, 8, app.footer_date)
    pdf.set_xy(pdf.w - pdf.r_margin - 80, y)
    pdf.cell(80, 8, app.footer_sign, align="R")
    return bytes(pdf.output())


def render_docx(app: LeaveApplication) -> bytes:
    doc = DocxDocument()
    section = doc.sections[0]
    section.left_margin, section.right_margin = Cm(2.5), Cm(1.5)
    section.top_margin, section.bottom_margin = Cm(2), Cm(2)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(14)

    for line in app.header:
        p = doc.add_paragraph(line)
        p.paragraph_format.left_indent = Cm(9)
        p.paragraph_format.space_after = Pt(0)

    doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(app.title).bold = True

    for paragraph in app.body:
        p = doc.add_paragraph(paragraph)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(1.25)

    doc.add_paragraph()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = app.footer_date
    right = table.cell(0, 1).paragraphs[0]
    right.text = app.footer_sign
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render(app: LeaveApplication, fmt: DocFormat) -> bytes:
    return render_pdf(app) if fmt == DocFormat.PDF else render_docx(app)


def filename(donation: Donation, fmt: DocFormat) -> str:
    return f"zayavlenie_den_otdyha_{donation.donated_on.isoformat()}.{fmt.value}"
